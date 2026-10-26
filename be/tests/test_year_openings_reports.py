"""
be/tests/test_year_openings_reports.py
BI-1b (D-043): reports follow the yearly openings. Point-in-time balances, cash flow (first-day rule, fiscal-year
adjustments, real reconciliation flag), yearly summary fields/warnings, the month view and its xlsx export,
balance sheet and trial balance regression.
"""
import io
import uuid
from datetime import date

import openpyxl
import pytest

from db import get_session_factory
from finance import year_openings
from finance.models import FinanceBankAccountDB, LedgerTransactionDB


@pytest.fixture(autouse=True)
def frozen_today(monkeypatch):
    monkeypatch.setattr(year_openings, "_today", lambda: date(2026, 10, 26))


@pytest.fixture
def db_session(app_client):
    session = get_session_factory()()
    try:
        yield session
    finally:
        session.close()


def make_account(db, opening=0.0, currency="USD", opening_date=None, name="BIR"):
    acc = FinanceBankAccountDB(account_name=f"{name} {uuid.uuid4().hex[:6]}", bank_name="Test Bank",
                               account_number=f"BR-{uuid.uuid4().hex[:8]}", currency=currency,
                               opening_balance=opening, opening_balance_date=opening_date,
                               current_balance=opening, is_active=True)
    db.add(acc)
    db.flush()
    return acc


def add_tx(db, acc, day, amount, direction):
    db.add(LedgerTransactionDB(account_id=acc.id, date=day, amount=amount, direction=direction,
                               currency=acc.currency, source="manual", entry_type="standard"))
    db.flush()


def put(client, cookies, acc, year, amount, note=None):
    body = {"opening_balance": amount}
    if note is not None:
        body["note"] = note
    r = client.put(f"/api/finance/accounts/{acc.id}/year-openings/{year}", json=body, cookies=cookies)
    assert r.status_code == 200, r.text


def summary(client, cookies, year, month=None):
    params = {"year": year}
    if month:
        params["month"] = month
    r = client.get("/api/finance/reports/fiscal-year-summary", params=params, cookies=cookies)
    assert r.status_code == 200, r.text
    return r.json()


def row(data, acc):
    return [a for a in data["accounts"] if a["account_id"] == acc.id][0]


def cash_flow(client, cookies, d_from, d_to, currency="USD"):
    r = client.get("/api/finance/reports/cash-flow", params={"date_from": d_from, "date_to": d_to, "currency": currency}, cookies=cookies)
    assert r.status_code == 200, r.text
    return r.json()


# 8. cash flow --------------------------------------------------------------------------------
def test_cash_flow_first_day_is_counted_once(app_client, admin_cookies, db_session):
    acc2 = make_account(db_session, 1000.0, currency="DKK", opening_date="2026-05-01")
    add_tx(db_session, acc2, "2026-05-01", 100.0, "out")
    db_session.commit()
    cf = cash_flow(app_client, admin_cookies, "2026-05-01", "2026-05-31", currency="DKK")
    assert cf["beginning_cash_balance"] == 1000.0
    assert cf["net_cash_operating"] == -100.0
    assert cf["ending_cash_balance"] == 900.0
    assert cf["is_reconciled"] is True and cf["unexplained_difference"] == 0.0
    assert cf["fiscal_year_adjustments"] == []


def test_cash_flow_range_crossing_january_first_lists_the_adjustment(app_client, admin_cookies, db_session):
    acc = make_account(db_session, 1000.0, currency="SEK")
    add_tx(db_session, acc, "2025-12-20", 200.0, "out")      # closing 2025: 800
    add_tx(db_session, acc, "2026-02-01", 30.0, "in")
    db_session.commit()
    put(app_client, admin_cookies, acc, 2026, 850.0)          # +50 against the previous closing
    cf = cash_flow(app_client, admin_cookies, "2025-12-01", "2026-03-31", currency="SEK")
    assert cf["beginning_cash_balance"] == 1000.0
    assert cf["fiscal_year_adjustments_total"] == 50.0
    assert cf["fiscal_year_adjustments"][0]["year"] == 2026 and cf["fiscal_year_adjustments"][0]["amount"] == 50.0
    assert cf["ending_cash_balance"] == 880.0                 # 850 + 30
    assert cf["net_change_in_cash"] == -170.0                 # -200 + 30
    assert cf["is_reconciled"] is True and cf["unexplained_difference"] == 0.0


def test_cash_flow_range_starting_on_january_first_uses_the_entered_opening(app_client, admin_cookies, db_session):
    acc = make_account(db_session, 1000.0, currency="NOK")
    add_tx(db_session, acc, "2026-01-01", 10.0, "out")
    db_session.commit()
    put(app_client, admin_cookies, acc, 2026, 700.0)
    cf = cash_flow(app_client, admin_cookies, "2026-01-01", "2026-01-31", currency="NOK")
    assert cf["beginning_cash_balance"] == 700.0 and cf["ending_cash_balance"] == 690.0
    assert cf["fiscal_year_adjustments"] == [] and cf["is_reconciled"] is True


def test_cash_flow_flags_a_real_mismatch(app_client, admin_cookies, db_session):
    acc = make_account(db_session, 500.0, currency="PLN")
    add_tx(db_session, acc, "2026-03-01", 100.0, "out")
    db_session.commit()
    # a movement the statement cannot see: a transaction in another currency on this account
    db_session.add(LedgerTransactionDB(account_id=acc.id, date="2026-03-02", amount=40.0, direction="out",
                                       currency="EUR", source="manual", entry_type="standard"))
    db_session.commit()
    cf = cash_flow(app_client, admin_cookies, "2026-03-01", "2026-03-31", currency="PLN")
    assert cf["is_reconciled"] is False and cf["unexplained_difference"] == -40.0


# 9. month view -------------------------------------------------------------------------------
@pytest.fixture
def month_account(app_client, admin_cookies, db_session):
    acc = make_account(db_session, 100.0, currency="JPY", opening_date="2025-01-01")
    for day, amt, d in [("2025-12-15", 20.0, "out"), ("2026-01-01", 5.0, "in"), ("2026-01-31", 7.0, "out"),
                        ("2026-02-10", 50.0, "in"), ("2026-03-31", 1.0, "out"), ("2026-06-30", 9.0, "in"),
                        ("2026-12-31", 3.0, "out")]:
        add_tx(db_session, acc, day, amt, d)
    db_session.commit()
    put(app_client, admin_cookies, acc, 2026, 90.0, note="bank statement")      # previous closing is 80: variance +10
    return acc


def test_month_view_invariants(app_client, admin_cookies, month_account):
    acc = month_account
    year = row(summary(app_client, admin_cookies, 2026), acc)
    assert year["opening"] == 90.0 and year["opening_source"] == "entered"
    assert year["previous_year_closing"] == 80.0 and year["opening_variance"] == 10.0 and year["opening_note"] == "bank statement"
    months = [row(summary(app_client, admin_cookies, 2026, m), acc) for m in range(1, 13)]
    for r in months:
        assert round(r["opening"] + r["money_in"] - r["money_out"], 2) == r["closing"]
    for prev, nxt in zip(months, months[1:]):
        assert nxt["opening"] == prev["closing"]
    assert months[0]["opening"] == year["opening"]
    assert months[-1]["closing"] == year["closing"]
    assert round(sum(r["money_in"] for r in months), 2) == year["money_in"]
    assert round(sum(r["money_out"] for r in months), 2) == year["money_out"]
    assert months[0]["opening_source"] == "entered" and months[1]["opening_source"] == "previous_month"
    jan = summary(app_client, admin_cookies, 2026, 1)
    assert jan["period"] == {"from": "2026-01-01", "to": "2026-01-31"} and jan["month"] == 1


def test_month_view_multi_currency_totals_stay_separate_and_bad_month_rejected(app_client, admin_cookies, db_session, month_account):
    other = make_account(db_session, 10.0, currency="KRW")
    add_tx(db_session, other, "2026-02-02", 4.0, "in")
    db_session.commit()
    totals = summary(app_client, admin_cookies, 2026, 2)["currency_totals"]
    assert totals["KRW"] == {"opening": 10.0, "money_in": 4.0, "money_out": 0.0, "closing": 14.0}
    assert totals["JPY"] == {"opening": 88.0, "money_in": 50.0, "money_out": 0.0, "closing": 138.0}
    assert app_client.get("/api/finance/reports/fiscal-year-summary", params={"year": 2026, "month": 13}, cookies=admin_cookies).status_code == 422


def test_yearly_summary_warnings(app_client, admin_cookies, db_session, month_account):
    data = summary(app_client, admin_cookies, 2026)
    mine = [w for w in data["warnings"] if w["account_id"] == month_account.id]
    assert [w["type"] for w in mine] == ["opening_variance", "transactions_before_first_base"]
    assert mine[0]["amount"] == 10.0 and mine[1]["amount"] == 1     # the 2025 transaction precedes the first entered opening
    carried = [w for w in summary(app_client, admin_cookies, 2027)["warnings"] if w["account_id"] == month_account.id]
    assert [w["type"] for w in carried] == ["carried_forward", "transactions_before_first_base"]
    # legacy accounts with no rows produce no variance / carried-forward warning
    legacy = make_account(db_session, 5.0, currency="CZK")
    db_session.commit()
    assert not [w for w in summary(app_client, admin_cookies, 2026)["warnings"] if w["account_id"] == legacy.id]


def test_no_data_rule_applies_only_without_an_entered_row(app_client, admin_cookies, db_session):
    late = make_account(db_session, 0.0, currency="HUF", opening_date="2030-01-01")
    entered = make_account(db_session, 0.0, currency="RON", opening_date="2030-01-01")
    db_session.commit()
    put(app_client, admin_cookies, entered, 2026, 12.0)
    data = summary(app_client, admin_cookies, 2026)
    assert row(data, late)["no_data"] is True
    assert row(data, entered)["no_data"] is False and row(data, entered)["opening"] == 12.0


def test_xlsx_carries_the_month(app_client, admin_cookies, month_account):
    resp = app_client.get("/api/finance/reports/fiscal-year-summary", params={"year": 2026, "month": 2, "format": "xlsx"}, cookies=admin_cookies)
    assert resp.status_code == 200 and "fiscal_year_2026_02_summary.xlsx" in resp.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Accounts"]
    assert "2026-02" in ws["A1"].value
    assert ws["A2"].value.startswith("2026-02-01 to 2026-02-28")
    headers = [c.value for c in ws[4]]
    assert "Opening source" in headers and "Difference from previous closing" in headers
    whole = app_client.get("/api/finance/reports/fiscal-year-summary", params={"year": 2026, "format": "xlsx"}, cookies=admin_cookies)
    assert "fiscal_year_2026_summary.xlsx" in whole.headers["content-disposition"]


# 10. point-in-time, balance sheet, trial balance ---------------------------------------------
def test_point_in_time_keys_and_base(app_client, admin_cookies, db_session):
    acc = make_account(db_session, 100.0, currency="THB")
    add_tx(db_session, acc, "2025-06-01", 30.0, "out")
    add_tx(db_session, acc, "2026-06-01", 5.0, "in")
    db_session.commit()
    put(app_client, admin_cookies, acc, 2026, 500.0)

    def at(d):
        rows = app_client.get("/api/finance/reports/balances", params={"as_of_date": d}, cookies=admin_cookies).json()["accounts"]
        return [r for r in rows if r["account_id"] == acc.id][0]

    r = at("2026-12-31")
    assert r["balance_as_of_date"] == 505.0 and r["opening_balance"] == 100.0       # legacy value kept
    assert r["base_year"] == 2026 and r["base_amount"] == 500.0 and r["base_source"] == "entered"
    old = at("2025-12-31")
    assert old["balance_as_of_date"] == 70.0 and old["base_year"] is None and old["base_source"] == "account_opening"


def test_balance_sheet_and_trial_balance_still_balance_and_use_entered_openings(app_client, admin_cookies, db_session):
    acc = make_account(db_session, 100.0, currency="USD")
    add_tx(db_session, acc, "2026-02-01", 30.0, "out")
    db_session.commit()
    put(app_client, admin_cookies, acc, 2026, 400.0)
    for path in ("balance-sheet", "trial-balance"):
        r = app_client.get(f"/api/finance/reports/{path}", params={"as_of_date": "2026-12-31", "currency": "USD"}, cookies=admin_cookies)
        assert r.status_code == 200, r.text
        assert r.json()["is_balanced"] is True
    bs = app_client.get("/api/finance/reports/balance-sheet", params={"as_of_date": "2026-12-31", "currency": "USD"}, cookies=admin_cookies).json()
    assert "370" in str(bs)
