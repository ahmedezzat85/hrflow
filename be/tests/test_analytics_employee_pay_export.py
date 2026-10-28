"""
be/tests/test_analytics_employee_pay_export.py
Export tests for Employee Pay Analytics (AC 7):
- XLSX reading back with openpyxl (sheets match endpoint row for row)
- CSV export for monthly and raises
- Formula injection neutralization (=, +, -, @)
- HR AuditRepository event logging (analytics.employee_pay.export)
"""
import io
import csv
import json
import pytest
import openpyxl
from db import get_db_context
from models_db import EmployeeDB, SalaryHistoryDB, AuditLogDB
from finance.models import PayrollRunDB, PayrollLineDB
from conftest import create_test_user


@pytest.fixture
def export_env(app_client):
    with get_db_context() as db:
        emp = EmployeeDB(
            name="Evan Export",
            email="evan.export@voyance.health",
            dept="Security",
            job_role="Security Engineer",
            status="active",
        )
        db.add(emp)
        db.commit()
        db.refresh(emp)
        emp_id = emp.id

        # Seed salary history with a formula injection attack in the reason
        h1 = SalaryHistoryDB(
            employee_id=emp_id,
            date="2026-01-01",
            reason="=1+1",        # starts with '='
            applied_by="admin@voyance.health",
        )
        h2 = SalaryHistoryDB(
            employee_id=emp_id,
            date="2026-03-01",
            reason="-cmd|' /C calc'!A0",  # starts with '-'
            applied_by="admin@voyance.health",
        )
        db.add_all([h1, h2])

        # Seed 2 payroll runs (one with commission, one with bonus)
        r1 = PayrollRunDB(
            period_label="2026-06",
            period_start="2026-06-01",
            period_end="2026-06-30",
            status="paid",
            currency="USD",
        )
        db.add(r1)
        db.flush()
        db.add(PayrollLineDB(payroll_run_id=r1.id, employee_id=emp_id, compensation_type="external_usd", base_salary=9000.0, payment_status="paid"))
        db.add(PayrollLineDB(payroll_run_id=r1.id, employee_id=emp_id, compensation_type="internal_usd_cash", base_salary=3500.0, payment_status="paid"))
        db.add(PayrollLineDB(payroll_run_id=r1.id, employee_id=emp_id, compensation_type="commission_sales", base_salary=1200.0, payment_status="paid"))

        r2 = PayrollRunDB(
            period_label="2026-07",
            period_start="2026-07-01",
            period_end="2026-07-31",
            status="paid",
            currency="USD",
        )
        db.add(r2)
        db.flush()
        db.add(PayrollLineDB(payroll_run_id=r2.id, employee_id=emp_id, compensation_type="external_usd", base_salary=9000.0, payment_status="paid"))
        db.add(PayrollLineDB(payroll_run_id=r2.id, employee_id=emp_id, compensation_type="internal_usd_cash", base_salary=3500.0, payment_status="paid"))
        db.add(PayrollLineDB(payroll_run_id=r2.id, employee_id=emp_id, compensation_type="bonus", base_salary=800.0, payment_status="paid"))

        db.commit()

    admin_cookies = create_test_user("admin_export@voyance.health", role_key="super_admin")
    finance_cookies = create_test_user("fin_export@voyance.health", role_key="financial_admin")
    hr_cookies = create_test_user("hr_export@voyance.health", role_key="hr_admin")

    return {
        "emp_id": emp_id,
        "admin_cookies": admin_cookies,
        "finance_cookies": finance_cookies,
        "hr_cookies": hr_cookies,
    }


def test_export_xlsx_matches_endpoint_and_audits(app_client, export_env):
    """AC 7: XLSX sheets match endpoint data row for row, and audit event is logged."""
    emp_id = export_env["emp_id"]
    cookies = export_env["admin_cookies"]

    # 1. Fetch JSON endpoint data
    json_resp = app_client.get(f"/api/analytics/employees/{emp_id}/pay", cookies=cookies)
    assert json_resp.status_code == 200
    pay_data = json_resp.json()
    series = pay_data["payments"]["series"]
    changes = pay_data["plan"]["changes"]

    # 2. Fetch XLSX export
    resp = app_client.get(
        f"/api/analytics/employees/{emp_id}/pay/export?format=xlsx",
        cookies=cookies,
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in resp.headers["content-disposition"]

    # Read back with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Monthly" in wb.sheetnames
    assert "Raises" in wb.sheetnames

    # Check Monthly sheet rows match series
    ws_m = wb["Monthly"]
    rows_m = list(ws_m.iter_rows(values_only=True))
    assert rows_m[0] == (
        "Period", "Base External", "Base Internal", "Base Total",
        "Commission", "Bonus", "Total", "Payment Status", "Paid Date"
    )
    assert len(rows_m) - 1 == len(series)
    for idx, s in enumerate(series):
        row = rows_m[idx + 1]
        assert row[0] == s["bucket"]
        assert row[1] == s["base_external"]
        assert row[2] == s["base_internal"]
        assert row[3] == round(s["base_external"] + s["base_internal"], 2)
        assert row[4] == s["commission"]
        assert row[5] == s["bonus"]
        assert row[6] == s["total"]
        assert row[7] == s["payment_status"]

    # Check Raises sheet rows match changes
    ws_r = wb["Raises"]
    rows_r = list(ws_r.iter_rows(values_only=True))
    assert rows_r[0] == ("Date", "Previous", "New", "Delta", "%", "Reason")
    assert len(rows_r) - 1 == len(changes)
    for idx, ch in enumerate(changes):
        row = rows_r[idx + 1]
        assert row[0] == ch["date"]
        assert row[1] == ch["previous_total"]
        assert row[2] == ch["new_total"]
        assert row[3] == round(ch["new_total"] - ch["previous_total"], 2)
        # Note: formula injection test checks neutralisation below

    # Check audit event logged via HR AuditRepository (Decision 9)
    with get_db_context() as db:
        audit = (
            db.query(AuditLogDB)
            .filter(
                AuditLogDB.action == "analytics.employee_pay.export",
                AuditLogDB.target_id == str(emp_id),
            )
            .order_by(AuditLogDB.id.desc())
            .first()
        )
        assert audit is not None
        assert audit.actor_email == "admin_export@voyance.health"
        details = json.loads(audit.details)
        assert details["format"] == "xlsx"
        assert "filters" in details


def test_export_formula_injection_neutralization(app_client, export_env):
    """AC 7: Formula triggers (=, +, -, @) are prefixed with single quote in XLSX and CSV."""
    emp_id = export_env["emp_id"]
    cookies = export_env["admin_cookies"]

    # Check XLSX
    resp_xlsx = app_client.get(f"/api/analytics/employees/{emp_id}/pay/export?format=xlsx", cookies=cookies)
    wb = openpyxl.load_workbook(io.BytesIO(resp_xlsx.content))
    ws_r = wb["Raises"]
    rows_r = list(ws_r.iter_rows(values_only=True))

    # Row 1: reason="=1+1"; row 2: reason="-cmd|' /C calc'!A0". The percentage is derived from the plan and always
    # starts with '+' or '-', so it must be neutralized too.
    # Both rows must be neutralized with leading single-quote
    reasons = [r[5] for r in rows_r[1:]]
    pcts = [r[4] for r in rows_r[1:]]

    for reason in reasons:
        assert str(reason).startswith("'"), f"Reason not neutralized: {reason}"
    for pct in pcts:
        assert str(pct).startswith("'"), f"Pct change not neutralized: {pct}"

    # Check CSV
    resp_csv = app_client.get(f"/api/analytics/employees/{emp_id}/pay/export?format=csv&dataset=raises", cookies=cookies)
    reader = csv.reader(io.StringIO(resp_csv.text))
    csv_rows = list(reader)
    csv_reasons = [r[5] for r in csv_rows[1:]]
    csv_pcts = [r[4] for r in csv_rows[1:]]

    for reason in csv_reasons:
        assert reason.startswith("'"), f"CSV Reason not neutralized: {reason}"
    for pct in csv_pcts:
        assert pct.startswith("'"), f"CSV Pct not neutralized: {pct}"


def test_export_csv_monthly_and_raises(app_client, export_env):
    """AC 7: CSV export generates accurate rows for dataset=monthly and dataset=raises."""
    emp_id = export_env["emp_id"]
    cookies = export_env["admin_cookies"]

    # Monthly CSV
    r_m = app_client.get(f"/api/analytics/employees/{emp_id}/pay/export?format=csv&dataset=monthly", cookies=cookies)
    assert r_m.status_code == 200
    assert "text/csv" in r_m.headers["content-type"]
    reader_m = csv.reader(io.StringIO(r_m.text))
    rows_m = list(reader_m)
    assert rows_m[0] == [
        "Period", "Base External", "Base Internal", "Base Total",
        "Commission", "Bonus", "Total", "Payment Status", "Paid Date"
    ]
    assert len(rows_m) == 3  # Header + 2 months

    # Raises CSV
    r_r = app_client.get(f"/api/analytics/employees/{emp_id}/pay/export?format=csv&dataset=raises", cookies=cookies)
    assert r_r.status_code == 200
    reader_r = csv.reader(io.StringIO(r_r.text))
    rows_r = list(reader_r)
    assert rows_r[0] == ["Date", "Previous", "New", "Delta", "%", "Reason"]
    assert len(rows_r) == 3  # Header + 2 raises


def test_export_permission_scoping(app_client, export_env):
    """AC 1 & AC 7: Export includes only sheets/datasets permitted by caller's role."""
    emp_id = export_env["emp_id"]

    # 1. Finance-only caller with format=xlsx gets Monthly sheet ONLY
    resp_fin_xlsx = app_client.get(
        f"/api/analytics/employees/{emp_id}/pay/export?format=xlsx",
        cookies=export_env["finance_cookies"],
    )
    assert resp_fin_xlsx.status_code == 200
    wb_fin = openpyxl.load_workbook(io.BytesIO(resp_fin_xlsx.content))
    assert "Monthly" in wb_fin.sheetnames
    assert "Raises" not in wb_fin.sheetnames

    # 2. HR-only caller with format=xlsx gets Raises sheet ONLY
    resp_hr_xlsx = app_client.get(
        f"/api/analytics/employees/{emp_id}/pay/export?format=xlsx",
        cookies=export_env["hr_cookies"],
    )
    assert resp_hr_xlsx.status_code == 200
    wb_hr = openpyxl.load_workbook(io.BytesIO(resp_hr_xlsx.content))
    assert "Raises" in wb_hr.sheetnames
    assert "Monthly" not in wb_hr.sheetnames

    # 3. Finance-only caller trying to export dataset=raises gets 403
    resp_fin_raises = app_client.get(
        f"/api/analytics/employees/{emp_id}/pay/export?format=csv&dataset=raises",
        cookies=export_env["finance_cookies"],
    )
    assert resp_fin_raises.status_code == 403

    # 4. HR-only caller trying to export dataset=monthly gets 403
    resp_hr_monthly = app_client.get(
        f"/api/analytics/employees/{emp_id}/pay/export?format=csv&dataset=monthly",
        cookies=export_env["hr_cookies"],
    )
    assert resp_hr_monthly.status_code == 403
