"""
be/finance/import_manifest.py
Manifest parsing for the bulk / historical import (plan 25, section 5). Pure functions: no database, no storage.

A manifest is a CSV (UTF-8, BOM accepted) or the first sheet of an XLSX, one row per document, one document
kind per run. Header names are exact and lowercase; dates are ISO; amounts are plain decimals.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Optional

from finance import import_config as cfg

COMMON_COLUMNS = ("row_ref", "source_file", "confidence", "review_flags", "notes")

KIND_COLUMNS: Dict[str, tuple] = {
    "bills": ("vendor_name", "bill_number", "issue_date", "due_date", "currency", "category", "total_amount", "description"),
    "cheques": ("account", "cheque_number", "payee", "amount", "currency", "issue_date"),
    "invoices": ("customer_name", "invoice_number", "issue_date", "due_date", "currency", "description", "amount", "vat_rate"),
}

# Headers that must be present. Cells may still be blank where a Draft allows it; whether a blank cell is
# acceptable is decided per kind by the importer. Optional headers are treated as blank when absent.
OPTIONAL_HEADERS = ("source_file", "confidence", "review_flags", "notes", "vat_rate")

_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_XLSX_MAGIC = b"PK\x03\x04"


class ManifestError(ValueError):
    """The manifest as a whole cannot be used (wrong format, missing columns, over the document limit)."""


@dataclass
class ManifestRow:
    number: int  # 1-based position among the data rows
    line: int  # row number as the user sees it in Excel / a text editor (header is line 1)
    values: Dict[str, str]

    def get(self, name: str) -> str:
        return self.values.get(name, "")


@dataclass
class Manifest:
    kind: str
    headers: List[str]
    rows: List[ManifestRow]
    warnings: List[str] = field(default_factory=list)


def all_columns(kind: str) -> tuple:
    return COMMON_COLUMNS + KIND_COLUMNS[kind]


def required_headers(kind: str) -> tuple:
    return tuple(c for c in all_columns(kind) if c not in OPTIONAL_HEADERS)


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------
def parse_iso_date(text: str) -> str:
    """Strict YYYY-MM-DD; raises ValueError. No day/month guessing."""
    s = (text or "").strip()
    if not _ISO_DATE.match(s):
        raise ValueError(f"'{s}' is not an ISO date (YYYY-MM-DD)")
    try:
        datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"'{s}' is not a real calendar date")
    return s


def parse_amount(text: str) -> float:
    """Plain decimal with '.'; no thousands separators, no currency symbols. Raises ValueError."""
    s = (text or "").strip()
    if not re.match(r"^-?\d+(\.\d+)?$", s):
        raise ValueError(f"'{s}' is not a plain decimal number (use '.' and no thousands separator)")
    try:
        d = Decimal(s)
    except InvalidOperation:  # pragma: no cover - the regex already excludes this
        raise ValueError(f"'{s}' is not a number")
    return float(d)


def parse_confidence(text: str) -> Optional[float]:
    s = (text or "").strip()
    if not s:
        return None
    try:
        value = parse_amount(s)
    except ValueError:
        raise ValueError(f"confidence '{s}' is not a number between 0 and 1")
    if not 0.0 <= value <= 1.0:
        raise ValueError(f"confidence {s} is outside 0 to 1")
    return value


def _cell_to_text(value) -> str:
    """XLSX cells arrive typed. Identifiers and filenames must stay strings (Sheets/Excel auto-typing)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)), "f")
    return str(value).strip()


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def _read_csv(content: bytes) -> List[List[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ManifestError("The manifest CSV is not UTF-8. Save it as 'CSV UTF-8' and upload it again.")
    first_line = text.split("\n", 1)[0]
    if ";" in first_line and "," not in first_line:
        raise ManifestError("The manifest CSV must be comma-separated (found ';' separators).")
    try:
        return [list(r) for r in csv.reader(io.StringIO(text, newline=""))]
    except csv.Error as e:
        raise ManifestError(f"The manifest CSV could not be read: {e}")


def _read_xlsx(content: bytes) -> List[List[str]]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ManifestError("The manifest XLSX could not be read. Check that it is a valid .xlsx file.")
    try:
        ws = wb.worksheets[0]
        return [[_cell_to_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_manifest(filename: str, content: bytes, kind: str) -> Manifest:
    """Parse and structurally validate a manifest. Row content is validated by the kind's importer."""
    if kind not in KIND_COLUMNS:
        raise ManifestError(f"Unknown import kind '{kind}'")
    if not content:
        raise ManifestError("The manifest file is empty.")
    if len(content) > cfg.MAX_MANIFEST_BYTES:
        raise ManifestError(f"The manifest is larger than {cfg.MAX_MANIFEST_BYTES // (1024 * 1024)} MB. {cfg.LIMIT_HELP}")

    name = (filename or "").lower()
    if name.endswith(".xlsx") or (not name.endswith(".csv") and content[:4] == _XLSX_MAGIC):
        if content[:4] != _XLSX_MAGIC:
            raise ManifestError("The manifest has an .xlsx name but is not an Excel workbook.")
        table = _read_xlsx(content)
    elif name.endswith(".csv"):
        if content[:4] == _XLSX_MAGIC:
            raise ManifestError("The manifest has a .csv name but is an Excel workbook. Upload it as .xlsx.")
        table = _read_csv(content)
    else:
        raise ManifestError("The manifest must be a .csv or .xlsx file.")

    # Header: first non-empty row
    table = [r for r in table]
    header_idx = next((i for i, r in enumerate(table) if any((c or "").strip() for c in r)), None)
    if header_idx is None:
        raise ManifestError("The manifest has no header row.")
    raw_headers = [(c or "").strip().lstrip("﻿") for c in table[header_idx]]
    headers = [h for h in raw_headers if h]
    if len(set(headers)) != len(headers):
        dupes = sorted({h for h in headers if headers.count(h) > 1})
        raise ManifestError(f"The manifest has repeated column names: {', '.join(dupes)}.")

    missing = [h for h in required_headers(kind) if h not in headers]
    if missing:
        raise ManifestError(
            f"The manifest is missing required column(s) for '{kind}': {', '.join(missing)}. "
            "Column names are exact and lowercase; download the template for this kind."
        )
    known = set(all_columns(kind))
    warnings = []
    unknown = [h for h in headers if h not in known]
    if unknown:
        warnings.append(f"Ignored unknown column(s): {', '.join(unknown)}.")

    rows: List[ManifestRow] = []
    for offset, raw in enumerate(table[header_idx + 1:], start=1):
        cells = [(c or "").strip() for c in raw]
        if not any(cells):
            continue
        values = {col: "" for col in all_columns(kind)}
        for h_i, h in enumerate(raw_headers):
            if h in known and h_i < len(cells):
                values[h] = cells[h_i]
        rows.append(ManifestRow(number=len(rows) + 1, line=header_idx + 1 + offset, values=values))
        if len(rows) > cfg.MAX_DOCUMENTS_PER_RUN:
            raise ManifestError(
                f"The manifest has more than {cfg.MAX_DOCUMENTS_PER_RUN} documents. {cfg.LIMIT_HELP}"
            )
    if not rows:
        raise ManifestError("The manifest has a header but no document rows.")
    return Manifest(kind=kind, headers=headers, rows=rows, warnings=warnings)
